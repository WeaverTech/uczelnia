import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).parent

PROGRAM = """\
AC JOINT SPEED9 ACCU1 TIMER0  OX= WX= #[-86.688,-7.0655,-86.28,0.95457,-100.13,-183.58]
AC JOINT SPEED9 ACCU1 TIMER0  OX= WX= #[-115.87,-13.166,-118.3,-1.0621,-74.347,-65.354]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-115.87,-8.7489,-128.17,-1.181,-60.069,-65.05]
AC LINEAR SPEED3 ACCU1 TIMER1  OX=13 WX= #[-115.87,-8.7492,-128.16,-1.181,-60.069,-65.05]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-115.87,-14.028,-113.48,-1.0387,-79.999,-65.46]
AC JOINT SPEED9 ACCU1 TIMER0  OX=13 WX= #[-63.174,-13.609,-113.17,-0.17866,-79.296,-118.25]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-63.179,-6.5709,-126.18,-0.20307,-59.272,-118.17]
AC LINEAR SPEED3 ACCU1 TIMER1  OX= WX= #[-63.179,-6.5709,-126.18,-0.20307,-59.272,-118.17]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-63.18,-11.698,-112.82,-0.17918,-77.742,-118.24]
AC JOINT SPEED9 ACCU1 TIMER0  OX= WX= #[-113.35,-10.567,-111.94,-0.99145,-78.085,-68.048]
AC JOINT SPEED9 ACCU1 TIMER0  OX= WX= #[-113.33,-10.549,-111.91,-0.99041,-78.084,-68.057]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-113.74,-4.8209,-124.9,-1.1317,-59.375,-67.282]
AC LINEAR SPEED3 ACCU1 TIMER1  OX=13 WX= #[-113.74,-4.8212,-124.9,-1.1312,-59.375,-67.282]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-113.74,-9.8796,-112.72,-1.0003,-76.61,-67.627]
AC JOINT SPEED9 ACCU1 TIMER0  OX=13 WX= #[-65.614,-7.4757,-110.72,-0.25241,-75.68,-115.85]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-65.617,-2.4974,-122.6,-0.28564,-58.83,-115.76]
AC LINEAR SPEED3 ACCU1 TIMER1  OX= WX= #[-65.617,-2.4974,-122.6,-0.28564,-58.83,-115.76]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-65.617,-7.975,-107.47,-0.24825,-79.419,-115.86]
AC JOINT SPEED9 ACCU1 TIMER0  OX= WX= #[-110.96,-5.8703,-105.64,-0.95353,-79.651,-70.395]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-111.9,-0.060775,-120.21,-1.1015,-59.289,-69.059]
AC LINEAR SPEED3 ACCU1 TIMER1  OX=13 WX= #[-111.9,-0.060775,-120.21,-1.1015,-59.289,-69.059]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-111.9,-5.3305,-106.64,-0.96756,-78.119,-69.424]
AC JOINT SPEED9 ACCU1 TIMER0  OX=13 WX= #[-68.177,-3.6911,-105.15,-0.30019,-77.479,-113.28]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-67.63,1.6322,-118.18,-0.32823,-59.12,-113.73]
AC LINEAR SPEED3 ACCU1 TIMER1  OX= WX= #[-67.63,1.6313,-118.18,-0.32823,-59.12,-113.73]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-67.631,-3.0578,-107.19,-0.29188,-74.798,-113.82]
AC JOINT SPEED9 ACCU1 TIMER0  OX= WX= #[-110.9,-2.5803,-106.74,-0.96859,-75.258,-70.377]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-110.56,3.5247,-116.2,-1.0823,-59.69,-70.425]
AC LINEAR SPEED3 ACCU1 TIMER1  OX=13 WX= #[-110.56,3.5259,-116.2,-1.0823,-59.69,-70.425]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-110.56,-1.3533,-104.12,-0.95925,-76.646,-70.75]
AC JOINT SPEED9 ACCU1 TIMER0  OX=13 WX= #[-69.902,1.0921,-101.68,-0.33446,-76.172,-111.54]
AC LINEAR SPEED3 ACCU1 TIMER0  OX=13 WX= #[-69.357,5.1525,-114.17,-0.36407,-59.62,-111.98]
AC LINEAR SPEED3 ACCU1 TIMER1  OX= WX= #[-69.357,5.151,-114.17,-0.36407,-59.62,-111.98]
AC LINEAR SPEED3 ACCU1 TIMER0  OX= WX= #[-69.357,0.41047,-83.543,-0.31473,-94.984,-112.2]
"""

SPEED = {"9": "100%", "3": "25%"}
MOTION = {"JOINT": "Ruch osiowy", "LINEAR": "Ruch liniowy"}
ACCU = {"1": "dokł. 1mm"}
STEP_RE = re.compile(
    r"AC (JOINT|LINEAR) SPEED(\d) ACCU(\d) TIMER(\d)\s+OX=(\d*) WX= #\[([^\]]+)\]"
)

TNR = "Times New Roman"
CALIBRI = "Calibri"
RED = RGBColor(0xFF, 0x00, 0x00)


def parse_steps():
    steps = []
    gripper_closed = False
    for line in PROGRAM.strip().splitlines():
        motion, speed, accu, timer, ox, joints = STEP_RE.match(line).groups()
        parts = [f"{MOTION[motion]} {SPEED[speed]}", ACCU[accu]]
        closed = ox == "13"
        if closed != gripper_closed:
            parts.append("zamknij chwytak" if closed else "otwórz chwytak")
            gripper_closed = closed
        if timer == "1":
            parts.append("czekaj 0,3s")
        steps.append(
            {
                "code": line,
                "comment": ", ".join(parts),
                "coords": ", ".join(j.strip() for j in joints.split(",")),
            }
        )
    return steps


def font(run, name, size, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), name)
    rpr.rFonts.set(qn("w:cs"), name)
    if color is not None:
        run.font.color.rgb = color
    return run


def fmt(p, spacing=1.0, align=None, left=None, first=None, tabs=(), exact=None):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if exact is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(exact)
    else:
        pf.line_spacing = spacing
    if align is not None:
        p.alignment = align
    if left is not None:
        pf.left_indent = Pt(left)
    if first is not None:
        pf.first_line_indent = Pt(first)
    for t in tabs:
        pf.tab_stops.add_tab_stop(Pt(t))
    return p


def table_borders(table, size_eighths):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size_eighths))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def col_widths(table, widths_pt):
    table.autofit = False
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table._tbl.tblPr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths_pt) * 20)))
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_pt):
        gc.set(qn("w:w"), str(int(w * 20)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_pt):
            cell.width = Pt(w)


def new_table(doc, rows, widths_pt, border):
    t = doc.add_table(rows=rows, cols=len(widths_pt))
    table_borders(t, border)
    col_widths(t, widths_pt)
    return t


def cell_par(cell, first=True):
    return cell.paragraphs[0] if first else cell.add_paragraph()


def empty(doc_or_cell, spacing, bold=True, size=12):
    p = doc_or_cell.add_paragraph()
    fmt(p, spacing)
    font(p.add_run(), TNR, size, bold)
    return p


def section_heading(doc, number, text, spacing):
    p = doc.add_paragraph()
    fmt(
        p,
        spacing,
        WD_ALIGN_PARAGRAPH.JUSTIFY,
        left=21.3,
        first=-21.3,
        tabs=(21.3,),
    )
    font(p.add_run(f"{number}\t{text}"), TNR, 12, bold=True)
    return p


def add_page_field(paragraph):
    run = font(paragraph.add_run(), CALIBRI, 11)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = Cm(2.25)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(2.5)
    sec.footer_distance = Cm(1.25)

    footer = sec.footer
    p = footer.paragraphs[0]
    fmt(p, 1.0, WD_ALIGN_PARAGRAPH.RIGHT)
    add_page_field(p)
    p2 = footer.add_paragraph()
    fmt(p2, 1.0, WD_ALIGN_PARAGRAPH.LEFT)
    font(p2.add_run("Wzorzec sprawozdania L4 – Podstawy Robotyki"), CALIBRI, 11)


def header_table(doc):
    t = new_table(doc, 4, (241.0, 247.0), 12)
    top = t.cell(0, 0).merge(t.cell(0, 1))
    p = fmt(cell_par(top), 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    font(p.add_run("Laboratorium z przedmiotu"), CALIBRI, 12)
    p = fmt(top.add_paragraph(), 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    font(p.add_run("Podstawy Robotyki"), CALIBRI, 12, bold=True)
    p = fmt(top.add_paragraph(), 1.0)
    font(p.add_run("Laboratorium nr 4"), CALIBRI, 12)
    p = fmt(top.add_paragraph(), 1.0, tabs=(48,))
    font(p.add_run("Temat:\t"), CALIBRI, 12)
    font(
        p.add_run("Programowanie robotów Kawasaki z wykorzystaniem Teach-Pendanta"),
        CALIBRI,
        12,
        bold=True,
    )

    p = fmt(cell_par(t.cell(1, 0)), 1.0, tabs=(106,))
    font(p.add_run("Rok akademicki:\t"), CALIBRI, 11)
    font(p.add_run("2026/2027"), CALIBRI, 11)
    p = fmt(cell_par(t.cell(1, 1)), 1.0)
    font(p.add_run("Data realizacji zajęć:  "), CALIBRI, 12)
    font(p.add_run("12.10.2026"), CALIBRI, 12, color=RED)
    t.rows[1].height = Pt(19.4)
    t.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    p = fmt(cell_par(t.cell(2, 0)), 1.0, tabs=(106,))
    font(p.add_run("Grupa:\t"), CALIBRI, 11)
    font(p.add_run("GL04"), CALIBRI, 11)
    font(fmt(cell_par(t.cell(2, 1)), 1.0).add_run(), CALIBRI, 12)
    t.rows[2].height = Pt(18)
    t.rows[2].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    names = t.cell(3, 0).merge(t.cell(3, 1))
    rows = (
        ("Imiona i Nazwiska:", "1. ……………………………………., odpowiedzialny za 1,…."),
        ("", "2. ……………………………………., odpowiedzialny za 2,...."),
        ("", "3. ……………………………………., odpowiedzialny za 3,…."),
        ("", "4. ……………………………………., odpowiedzialny za ?,….."),
    )
    for i, (label, line) in enumerate(rows):
        p = fmt(cell_par(names, first=i == 0), 1.15, tabs=(99,))
        font(p.add_run(f"{label}\t"), CALIBRI, 12)
        font(p.add_run(line), CALIBRI, 12, color=RED)


def boxed(doc, title, page_break_before=False):
    t = new_table(doc, 2, (482.0,), 4)
    p = fmt(cell_par(t.cell(0, 0)), 1.5, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.page_break_before = page_break_before
    font(p.add_run(title), TNR, 12)
    return t.cell(1, 0)


def build():
    steps = parse_steps()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    setup_page(doc)
    header_table(doc)

    empty(doc, 1.15)
    section_heading(
        doc,
        "1)",
        "Rysunek podanego do zaprogramowania zadania manipulacyjnego z zaznaczonym "
        "globalnym układem współrzędnych i parametrami ruchu (może być odręczny)",
        1.15,
    )
    for _ in range(10):
        empty(doc, 1.15)

    section_heading(
        doc,
        "2)",
        "Kod programu z Teach Pendanta (zdjęcie lub przepisane) wraz komentarzem do "
        "każdego kroku.",
        1.5,
    )

    code_lines = [".PROGRAM m13x04()"] + [s["code"] for s in steps] + [".END"]
    code_cell = boxed(doc, "Kod programu")
    for i, line in enumerate(code_lines):
        p = fmt(cell_par(code_cell, first=i == 0), exact=7.4)
        font(p.add_run(f"{i:>2}  {line}"), "Courier New", 7)

    sep = doc.add_paragraph()
    fmt(sep, exact=1)
    font(sep.add_run(), TNR, 1)

    comments = ["Start programu"] + [s["comment"] for s in steps] + ["Koniec programu"]
    comment_cell = boxed(doc, "Komentarz do linii programu", page_break_before=True)
    for i, text in enumerate(comments):
        p = fmt(cell_par(comment_cell, first=i == 0), 1.0, left=21.3, first=-21.3, tabs=(21.3,))
        font(p.add_run(f"{i}\t{text}"), TNR, 10)

    empty(doc, 1.5)
    section_heading(
        doc,
        "3)",
        "Współrzędne zaprogramowanych punktów dla poszczególnych kroków(mogą być "
        "zdjęcia z Teach Pendanta)",
        1.5,
    )
    empty(doc, 1.5)

    coords = new_table(doc, len(steps) + 1, (117.0, 371.0), 4)
    for cell, title in zip(
        coords.rows[0].cells, ("Krok (step)", "Współrzędne (x, y, z, o, a, t)")
    ):
        p = fmt(cell_par(cell), 1.5, WD_ALIGN_PARAGRAPH.CENTER)
        font(p.add_run(title), TNR, 12)
    for tr in coords._tbl.tr_lst:
        tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, s in enumerate(steps, start=1):
        row = coords.rows[i].cells
        p = fmt(cell_par(row[0]), 1.5, WD_ALIGN_PARAGRAPH.CENTER)
        font(p.add_run(str(i)), TNR, 12)
        p = fmt(cell_par(row[1]), 1.5, WD_ALIGN_PARAGRAPH.CENTER)
        font(p.add_run(s["coords"]), TNR, 12)

    empty(doc, 1.15)

    out = OUT_DIR / "Sprawozdanie_L4_m13x04.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build())
